"""Roadmap v7 Wave 4 — standards currency (docs/roadmap-v7.md §2, §4)."""

from __future__ import annotations

import asyncio
import json
from datetime import date, timedelta
from pathlib import Path


def _ctx():
    from impact_vision.tools.base import ToolExecutionContext

    return ToolExecutionContext(cwd=Path.cwd())


# ------------------------------------------------------- W4.1 ESRS / VSME law
def test_revised_esrs_is_law_with_date_aware_status() -> None:
    from impact_vision.impact.frameworks.esrs import (
        legal_status,
        load_simplified_datapoints,
        simplified_esrs_metadata,
    )

    meta = simplified_esrs_metadata(today=date(2026, 10, 6))
    assert meta["status"] == "published_oj"
    assert meta["legal_status"] == "published_not_in_force"
    assert meta["legal_instrument"].endswith("2026/1563")
    assert simplified_esrs_metadata(today=date(2026, 12, 1))["legal_status"] == "in_force_not_yet_applicable"
    assert simplified_esrs_metadata(today=date(2027, 1, 1))["legal_status"] == "applicable"
    assert legal_status() == "draft"
    rows = load_simplified_datapoints()
    assert rows and not any(r.synthetic or "DRAFT" in r.datapoint_id for r in rows)


def test_vsme_questionnaire_template() -> None:
    from impact_vision.impact.investee_collection import generate_investee_questionnaire_schema

    basic = generate_investee_questionnaire_schema(sector="vsme")
    comp = generate_investee_questionnaire_schema(sector="vsme_comprehensive")
    assert basic.sector == "vsme" and comp.metric_count > basic.metric_count
    ids = {f.metric_id for s in basic.sections for f in s.fields}
    assert "VSME-B3-scope1_tco2e" in ids and "VSME-B1" in ids
    assert not any(i.startswith("VSME-C") for i in ids)


def test_vsme_value_chain_cap_is_enforced_on_request_packs() -> None:
    from impact_vision.impact.engagements.data_room import build_data_request_pack

    capped = build_data_request_pack(
        engagement_id="e", bundle_id="dd_full_iwa",
        counterparty_employees=250, requester_in_csrd_scope=True,
    )
    assert capped.vsme_ceiling_applied
    by_id = {f.metric_id: f for f in capped.fields}
    assert by_id["OI1479"].required  # total GHG sits inside VSME B3
    assert not by_id["PI4060"].required  # outside VSME → voluntary
    assert any("PI4060" in n for n in capped.ceiling_notes)

    large = build_data_request_pack(
        engagement_id="e", bundle_id="dd_full_iwa",
        counterparty_employees=5000, requester_in_csrd_scope=True,
    )
    assert not large.vsme_ceiling_applied and all(f.required for f in large.fields)

    vsme = build_data_request_pack(engagement_id="e", bundle_id="vsme_basic")
    assert vsme.fields and all("VSME" in f.frameworks for f in vsme.fields)


# ----------------------------------------------------- W4.2 SFDR 2.0 Parliament
def test_sfdr2_impact_language_requires_art7_or_9_toc_and_measurement() -> None:
    from impact_vision.impact.frameworks.sfdr_pai import SFDR2Input, classify_sfdr2_category

    bad = classify_sfdr2_category(SFDR2Input(
        current_article=8, description="Our impact fund delivers measurable impact.",
        position="parliament",
    ))
    assert bad.position_label.startswith("European Parliament")
    assert bad.impact_check.uses_impact_language and bad.impact_check.eligible is False
    assert bad.impact_check.category_ok is False
    assert any("PAI" in c for c in bad.caveats)

    good = classify_sfdr2_category(SFDR2Input(
        current_article=9, pct_strategy_aligned=85, description="An impact fund.",
        impact_objective_predefined=True,
        toc_validation={"is_passing": True, "severity_counts": {"medium": 1}},
        evidence_provenance="evidence-based",
    ))
    assert good.impact_check.eligible is True

    failing_toc = classify_sfdr2_category(SFDR2Input(
        current_article=9, description="An impact fund.", impact_objective_predefined=True,
        toc_validation={"is_passing": False, "severity_counts": {"critical": 1}},
        has_measurable_outcomes=True,
    ))
    assert any("theory of change" in f for f in failing_toc.impact_check.failures)

    plain = classify_sfdr2_category(SFDR2Input(current_article=8, description="ESG screened fund"))
    assert plain.impact_check.uses_impact_language is False


def test_sfdr2_council_flexibility_only_in_council_position() -> None:
    from impact_vision.impact.frameworks.sfdr_pai import SFDR2Input, classify_sfdr2_category

    kw = dict(current_article=8, has_transition_plan=True, pct_strategy_aligned=50, fund_in_ramp_up=True)
    council = classify_sfdr2_category(SFDR2Input(**kw))
    parliament = classify_sfdr2_category(SFDR2Input(position="parliament", **kw))
    assert any(c.startswith("Council position") for c in council.caveats)
    assert not any(c.startswith("Council position") for c in parliament.caveats)
    assert any("trilogue" in c for c in parliament.caveats)


# ------------------------------------------------ W4.3 / W4.4 UK SRS + Hong Kong
def test_uk_profile_has_uk_srs() -> None:
    from impact_vision.impact.engagements.regulatory import get_jurisdiction_profile

    uk = get_jurisdiction_profile("United Kingdom")
    ob = next(o for o in uk.obligations if o.obligation_id == "uk_srs_listed_issuers")
    assert "comply-or-explain" in ob.summary and "PS26/19" in ob.framework
    assert ob.source_url.startswith("https://www.fca.org.uk/")


def test_hong_kong_profile_and_calendar() -> None:
    from impact_vision.impact.regulatory_calendar import build_regulatory_calendar, issb_status

    cal = build_regulatory_calendar(jurisdiction="Hong Kong", fiscal_year_end="2026-12-31")  # type: ignore[arg-type]
    titles = " ".join(i.title for i in cal.items)
    assert "HKEX" in titles and "HKSSA 5000" in titles and "Taxonomy" in titles
    assert issb_status("hk")["source_quality"] == "deep_link"


def test_hk_taxonomy_screen_wraps_eu_alignment_maths() -> None:
    from impact_vision.impact.frameworks.hk_taxonomy import screen_hk_taxonomy

    res = screen_hk_taxonomy(
        "SunCo",
        text="We build rooftop solar PV and EV charging stations; we also run an evergreen fund.",
        activities=[{"name": "Solar", "revenue_share_pct": 70, "eligible": True,
                     "substantial_contribution": True, "dnsh_pass": True},
                    {"name": "Charging", "revenue_share_pct": 30, "eligible": True}],
    )
    assert {c.activity_id for c in res.candidates} >= {"HK-PG-1", "HK-TR-2"}
    assert res.alignment.revenue_eligible_pct == 100 and res.alignment.revenue_aligned_pct == 70
    assert any("Hong Kong Taxonomy" in r for r in res.alignment.references)
    assert any(p["phase"] == "2B" and p["status"] == "consultation" for p in res.phases)


def test_hk_taxonomy_via_framework_tool() -> None:
    from impact_vision.tools.impact.framework_tool import FrameworkInput, FrameworkTool

    out = asyncio.run(FrameworkTool().execute(
        FrameworkInput(framework="hk_taxonomy", action="assess", description="Wind farm operator"), _ctx()
    ))
    assert not out.is_error and "HK-PG-2" in out.output


# ------------------------------------------------------------ W4.5 AI provenance
def test_ai_provenance_rules_vs_llm() -> None:
    from impact_vision.impact.ai_provenance import ai_provenance_for_report

    rules = ai_provenance_for_report({
        "impact_claims": [{"text": "x", "extracted_by": "regex"}],
        "five_dimensions": {"what": {"provenance": "estimated"}, "who": {"provenance": "evidence-based"}},
    })
    assert not rules.ai_generated and rules.extraction == "rules"
    assert rules.estimated_figures == 1 and rules.total_figures == 2
    assert "no generative AI" in rules.disclosure

    llm = ai_provenance_for_report({"impact_claims": [{"text": "x", "extracted_by": "llm"}]})
    assert llm.ai_generated and llm.tagging == "llm" and "Extracted by AI" in llm.badges

    fallback = ai_provenance_for_report({"impact_claims": [{"extracted_by": "llm-fallback(no key)"}]})
    assert not fallback.ai_generated

    declared = ai_provenance_for_report({"ai_usage": {"drafting": "llm", "model": "claude-x"}})
    assert declared.ai_generated and "claude-x" in declared.disclosure


def test_ai_provenance_on_every_output(tmp_path) -> None:
    from openpyxl import load_workbook

    from impact_vision.impact.exports import to_csv, to_json
    from impact_vision.impact.pipeline import assess_file, sample_deck_path, write_deliverables

    bundle = assess_file(sample_deck_path("solar"))
    files = {p.name.split("_", 1)[-1]: p for p in write_deliverables(bundle, tmp_path, include_data=True)}
    needle = "Automated analysis"
    for key in ("impact_report.html", "ic_memo.html"):
        match = next(p for name, p in files.items() if name.endswith(key))
        assert needle in match.read_text(encoding="utf-8"), key
    dd = next(p for name, p in files.items() if name.endswith("dd_report.html"))
    assert "Machine-generated document" in dd.read_text(encoding="utf-8")
    summary = json.loads(next(p for n, p in files.items() if n.endswith("summary.json")).read_text())
    assert summary["ai_disclosure"].startswith(needle)
    assert json.loads(to_json(bundle.report_data))["ai_provenance"]["ai_generated"] is False
    assert needle in to_csv(bundle.report_data)
    wb = load_workbook(next(p for n, p in files.items() if n.endswith("data.xlsx")))
    assert "AI provenance" in wb.sheetnames
    docx = [p for n, p in files.items() if n.endswith(".docx")]
    if docx:
        from docx import Document

        for path in docx:
            text = " ".join(par.text for par in Document(str(path)).paragraphs)
            assert "disclosure" in text.lower(), path.name


def test_ai_badge_localised_in_report() -> None:
    from impact_vision.impact.pipeline import assess_file, sample_deck_path
    from impact_vision.impact.report_templates.decision_report import render_decision_report

    data = assess_file(sample_deck_path("solar")).report_data
    assert "以規則擷取" in render_decision_report(data, lang="zh-HK")
    assert "以规则撷取" in render_decision_report(data, lang="zh-CN")


# ------------------------------------------------------------- W4.6 ISSA 5000
def test_issa_5000_is_default_for_new_periods() -> None:
    from impact_vision.impact.ai_provenance import AIUseRecord
    from impact_vision.impact.assurance import build_assurance_pack, recommended_assurance_standard

    assert recommended_assurance_standard("2026-12-15") == "ISSA5000"
    assert recommended_assurance_standard("2026-12-14") == "ISAE3000"
    assert recommended_assurance_standard("2026-12-14", climate_only=True) == "ISAE3410"
    assert recommended_assurance_standard(None, jurisdiction="Hong Kong") == "HKSSA5000"
    pack = build_assurance_pack(
        fund_name="F", reporting_period="FY2027", assertion_text="a", prepared_by="CFO",
        subject_description="s", metrics=[], jurisdiction="HK",
        ai_use=[AIUseRecord(figure="5D what", stage="calculation", method="deterministic")],
    )
    assert pack.standard == "HKSSA5000" and pack.ai_use[0].stage == "calculation"


def test_issa5000_pack_carries_ai_use_register() -> None:
    from impact_vision.impact.assurance import build_issa5000_pack
    from impact_vision.impact.audit_trail import AuditTrail
    from impact_vision.impact.evidence_graph import EvidenceGraph

    trail = AuditTrail()
    pack = build_issa5000_pack(
        {"assertions": [], "impact_claims": [{"text": "c", "extracted_by": "llm"}]},
        EvidenceGraph(), trail, "limited",
    )
    assert pack["jurisdictional_equivalents"]["HK"] == "HKSSA 5000"
    assert any(r["stage"] == "extraction" and r["method"] == "llm" for r in pack["ai_use_register"])


# ----------------------------------------------------------------- W4.7 ECGT
def test_ecgt_is_operative_and_gcd_is_best_practice() -> None:
    from impact_vision.impact.greenwashing import assess_green_claims_compliance

    res = assess_green_claims_compliance(
        "Eco-friendly packaging; our product is carbon neutral thanks to offsets. Net zero by 2030."
    )
    assert "2024/825" in res.operative_law and "shelved" in res.gcd_status
    joined = " ".join(res.ecgt_breaches)
    assert "point 4a" in joined and "point 4c" in joined and "Art 6(2)(d)" in joined
    assert all("GCD" in g for g in res.best_practice_gaps if "shelved" in g)
    assert not res.compliant

    ok = assess_green_claims_compliance(
        "EU Ecolabel certified detergent. Net zero by 2040 under our SBTi transition plan.",
        has_lca=True, has_independent_verification=True,
    )
    assert not ok.ecgt_breaches and ok.compliant


def test_greenwashing_tool_reports_ecgt() -> None:
    from impact_vision.tools.impact.greenwashing_tool import GreenwashingDetectorTool, GreenwashingInput

    out = asyncio.run(GreenwashingDetectorTool().execute(GreenwashingInput(
        company_name="Acme", company_description="Our eco-friendly bottles are climate neutral via offsets.",
    ), _ctx()))
    assert "ECGT" in out.output and out.metadata["eu_green_claims"]["ecgt_breaches"]


# ------------------------------------------------------- W4.8 / W4.9 ISSB + Asia
def test_registry_tracks_wave4_standards() -> None:
    from impact_vision.impact.standards_registry import default_standards_registry

    reg = default_standards_registry()
    assert reg.get("Human Capital project").name.endswith("Workforce-related Disclosures")
    assert reg.get("HKSSA 5000").effective_date == "2026-12-15"
    assert reg.get("ECGT").status == "active"
    assert reg.get("ESRS_SIMPLIFIED_2026").status == "active"
    assert reg.get("ISSB_NATURE").status == "draft"


def test_tnfd_feeds_issb() -> None:
    from impact_vision.impact.frameworks.tnfd import TNFD_STATUS, TNFDInput, assess_tnfd

    assert "ISSB" in TNFD_STATUS
    assert "ISSB" in assess_tnfd(TNFDInput(company_name="X")).status


def test_asia_adoption_rows_are_refreshed() -> None:
    from impact_vision.impact.regulatory_calendar import issb_status

    assert "¥3tn" in issb_status("Japan")["scope"]
    assert "2024-11-20" in issb_status("China")["scope"]
    assert "S$1bn" in issb_status("Singapore")["scope"]
    assert "2026-07-01" in issb_status("Australia")["scope"]
    for j in ("Japan", "China", "Singapore", "Australia", "Hong Kong", "United Kingdom"):
        assert issb_status(j)["last_verified"] == "2026-10-06"


# -------------------------------------------------------------- W4.10 SB 253
def test_statutory_deadline_overrides_fiscal_offset_and_alerts() -> None:
    from impact_vision.impact.engagements import regulatory
    from impact_vision.impact.regulatory_calendar import calendar_alerts, _calendar_item

    deadlines = regulatory.schedule_deadlines(
        engagement_id="e", jurisdiction="US", fiscal_year_end="2026-12-31"
    )
    sb253 = next(d for d in deadlines if d.obligation_id == "ca_sb253_ghg_report")
    if date.today() <= date(2026, 11, 10):
        assert sb253.due_date == "2026-11-10" and sb253.statutory
        alerts = calendar_alerts([_calendar_item(d) for d in deadlines])
        if date(2026, 11, 10) - date.today() <= timedelta(days=60):
            assert any("SB 253" in a for a in alerts)
    else:  # after the first deadline the fiscal-year rule takes over again
        assert not sb253.statutory


def test_watchlist_rows_carry_sources() -> None:
    from impact_vision.impact.regulatory_calendar import regulatory_watchlist

    rows = regulatory_watchlist(today=date(2026, 10, 6))
    events = " ".join(r.event for r in rows)
    for needle in ("2026/1563", "SB 253", "HKSSA 5000", "UK SRS", "XBRL", "Art 50"):
        assert needle in events, needle
    assert all(r.source_url for r in rows if "SB 253" in r.event or "2026/1563" in r.event)
