"""v8 W4: machine-readable AI Act Art 50 marking on every output format."""
from __future__ import annotations

from pathlib import Path

import pytest

from impact_vision.impact.ai_provenance import (
    AIProvenance,
    IPTC_DST,
    machine_marking,
    mark_docx,
    mark_html,
    mark_pdf,
    mark_xlsx,
    marking_from_html,
)
from impact_vision.impact.pipeline import assess_file, write_deliverables

DECK = Path(__file__).parent / "golden_decks" / "loopwear_hk.md"


def test_rules_engine_is_algorithmic_and_llm_is_composite():
    assert machine_marking(AIProvenance())["digital_source_type"] == IPTC_DST + "algorithmicMedia"
    llm = AIProvenance(extraction="llm")
    assert machine_marking(llm)["digital_source_type"] == IPTC_DST + "compositeWithTrainedAlgorithmicMedia"


def test_html_marking_round_trips_and_is_idempotent():
    marking = machine_marking()
    doc = mark_html(mark_html("<html><head><title>x</title></head><body></body></html>", marking), marking)
    assert doc.count('id="ai-marking"') == 1
    back = marking_from_html(doc)
    assert back["digital_source_type"] == marking["digital_source_type"]
    assert back["disclosure"] == marking["disclosure"]


def test_pdf_gets_info_and_xmp(tmp_path):
    import pymupdf

    path = tmp_path / "x.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "hello")
    doc.save(path)
    doc.close()
    mark_pdf(path, machine_marking())
    doc = pymupdf.open(path)
    assert "digitalsourcetype=" in doc.metadata["keywords"]
    assert "Iptc4xmpExt:DigitalSourceType" in doc.get_xml_metadata()


def test_docx_and_xlsx_properties():
    docx = pytest.importorskip("docx")
    openpyxl = pytest.importorskip("openpyxl")
    d = docx.Document()
    mark_docx(d)
    assert d.core_properties.keywords.startswith("digitalsourcetype=")
    wb = openpyxl.Workbook()
    mark_xlsx(wb)
    assert wb.properties.keywords.startswith("digitalsourcetype=")


def test_every_deliverable_is_marked(tmp_path):
    pytest.importorskip("docx")
    bundle = assess_file(DECK)
    files = write_deliverables(bundle, tmp_path)
    for f in files:
        if f.suffix == ".html":
            assert 'name="iptc:digitalsourcetype"' in f.read_text(encoding="utf-8"), f.name
        elif f.suffix == ".docx":
            import docx

            assert docx.Document(str(f)).core_properties.keywords.startswith("digitalsourcetype="), f.name
        elif f.suffix == ".xlsx":
            import openpyxl

            assert openpyxl.load_workbook(f).properties.keywords.startswith("digitalsourcetype="), f.name
    assert bundle.summary()["ai_marking"]["machine_generated"] is True
