"""Import reported metrics from Excel / CSV (roadmap v8 W3.5).

Investees and portfolio teams keep KPIs in spreadsheets. Drop one next to the
deck and its figures count as *reported* metrics (not estimates):

* a column naming the metric — an IRIS+ ID (``PI4060``) or a metric name
  ("Client Individuals: Total", matched against the IRIS+ catalogue);
* a value column;
* optionally a period / year and a unit column. With several periods the most
  recent value is used and the full series is kept for trends.

Headers are matched case-insensitively against common synonyms, so most
templates (ILPA, EDCI, a fund's own) work without mapping.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

SPREADSHEET_SUFFIXES = (".csv", ".xlsx")
_METRIC_HEADERS = ("iris id", "iris+ id", "iris", "metric id", "id", "metric", "indicator", "kpi", "metric name",
                   "indicator name", "name")
_VALUE_HEADERS = ("value", "amount", "result", "actual", "reported value", "figure", "total")
_PERIOD_HEADERS = ("period", "year", "fiscal year", "fy", "reporting period", "date")
_UNIT_HEADERS = ("unit", "units", "uom")
_IRIS_ID = re.compile(r"^(?:OI|PI|PD|OD|FP|PDIV)\d{4}$", re.I)


@dataclass
class ImportResult:
    metrics: dict[str, Any] = field(default_factory=dict)        # IRIS+ ID → latest value
    series: dict[str, list[tuple[str, Any]]] = field(default_factory=dict)
    unmatched: list[str] = field(default_factory=list)
    rows: int = 0
    source: str = ""


def _rows(path: Path) -> list[list[Any]]:
    if path.suffix.lower() == ".csv":
        raw = path.read_text(encoding="utf-8-sig", errors="replace")
        dialect = csv.Sniffer().sniff(raw[:2048], delimiters=",;\t") if raw.strip() else csv.excel
        return [row for row in csv.reader(io.StringIO(raw), dialect)]
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    out: list[list[Any]] = []
    for ws in wb.worksheets:
        out.extend([list(r) for r in ws.iter_rows(values_only=True)])
        if out:
            break  # first non-empty sheet
    return out


def _find(headers: list[str], options: tuple[str, ...]) -> int | None:
    for opt in options:
        if opt in headers:
            return headers.index(opt)
    for i, h in enumerate(headers):
        if any(h.startswith(opt) for opt in options):
            return i
    return None


def _number(value: Any) -> Any:
    if isinstance(value, (int, float)):
        return value
    text = str(value or "").strip().replace(",", "")
    try:
        return float(text.rstrip("%")) if text else None
    except ValueError:
        return str(value).strip() or None


def _metric_id(label: str, by_name: dict[str, str]) -> str | None:
    label = str(label or "").strip()
    if _IRIS_ID.match(label):
        return label.upper()
    head = label.split()[0] if label else ""
    if _IRIS_ID.match(head):
        return head.upper()
    return by_name.get(label.lower())


def import_metrics(path: str | Path) -> ImportResult:
    """Read a spreadsheet of reported metrics; unknown labels are listed, not guessed."""
    p = Path(path)
    rows = [r for r in _rows(p) if any(c not in (None, "") for c in r)]
    result = ImportResult(source=p.name)
    if not rows:
        return result
    headers = [str(h or "").strip().lower() for h in rows[0]]
    mi, vi = _find(headers, _METRIC_HEADERS), _find(headers, _VALUE_HEADERS)
    pi = _find(headers, _PERIOD_HEADERS)
    if mi is None or vi is None:
        # No header row: assume "metric, value[, period]".
        mi, vi, pi = 0, 1, 2 if len(rows[0]) > 2 else None
    else:
        rows = rows[1:]
    from openharness.impact.database import get_metric_store

    by_name = {m.name.lower(): m.id for m in get_metric_store().all_metrics()}
    latest: dict[str, tuple[str, Any]] = {}
    for row in rows:
        if mi >= len(row) or vi >= len(row):
            continue
        label, value = row[mi], _number(row[vi])
        if label in (None, "") or value is None:
            continue
        result.rows += 1
        metric_id = _metric_id(str(label), by_name)
        if metric_id is None:
            result.unmatched.append(str(label))
            continue
        period = str(row[pi]).strip() if pi is not None and pi < len(row) and row[pi] not in (None, "") else ""
        result.series.setdefault(metric_id, []).append((period, value))
        if metric_id not in latest or period >= latest[metric_id][0]:
            latest[metric_id] = (period, value)
    result.metrics = {k: v for k, (_, v) in latest.items()}
    return result


__all__ = ["ImportResult", "SPREADSHEET_SUFFIXES", "import_metrics"]
