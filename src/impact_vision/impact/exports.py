"""Machine-readable exports of an impact report (W2.8).

One place for the three data formats an analyst opens next to the HTML/PDF:

* :func:`to_json` – the report payload with ``schema_version`` and an
  optional ``slim`` mode that drops evidence chains, catalog definitions and
  duplicated keys (the pig-farm payload shrinks ~6x, 160 KB to 27 KB).
* :func:`to_csv` – one tidy row per figure. ``Value`` keeps the display
  string (``2.3/5.0``) and ``Numeric`` / ``Max`` / ``Unit`` carry the number,
  so a spreadsheet can sort and chart without parsing text.
* :func:`write_xlsx` – a workbook with frozen headers, autofilters, sized
  columns, real numeric cells and a Methodology sheet.
"""
from __future__ import annotations

import copy
import csv
import io
import json
import re
from pathlib import Path
from typing import Any

SCHEMA_VERSION = "1.0"

DIMENSIONS = ("what", "who", "how_much", "contribution", "risk")

# Keys that only matter for the full audit trail (or are duplicates).
_SLIM_DROP_TOP = ("sdg_alignment",)
_SLIM_DROP_SDG = ("evidence_chain", "provenance", "scoring_basis")
_SLIM_GAP_KEEP = (
    "core_metric_set_basis", "core_metric_set_size", "metrics_reported",
    "metrics_missing", "coverage_percentage", "missing_by_dimension", "recommendations",
)

_NUMBER = re.compile(r"^\s*([-+]?\d[\d,]*(?:\.\d+)?)\s*(%|[^\d\s].*)?$")


def parse_number(value: Any) -> tuple[float | None, str]:
    """Split ``"1,200 tCO2e"`` into ``(1200.0, "tCO2e")``; non-numbers give ``(None, "")``."""
    if isinstance(value, bool):
        return None, ""
    if isinstance(value, (int, float)):
        return float(value), ""
    match = _NUMBER.match(str(value or ""))
    if not match:
        return None, ""
    try:
        number = float(match.group(1).replace(",", ""))
    except ValueError:
        return None, ""
    return number, (match.group(2) or "").strip()


def _num(value: Any) -> float | None:
    return parse_number(value)[0]


# ---------------------------------------------------------------- JSON


def slim_report(data: dict) -> dict:
    """A compact copy of *data* for pipelines: scores, claims and gaps without
    evidence chains, IRIS+ definitions or duplicated keys."""
    out = {k: copy.deepcopy(v) for k, v in data.items() if k not in _SLIM_DROP_TOP}
    for alignment in out.get("sdg_alignments", []) or []:
        for key in _SLIM_DROP_SDG:
            alignment.pop(key, None)
    gap = out.get("gap_analysis")
    if isinstance(gap, dict):
        slim_gap = {k: gap[k] for k in _SLIM_GAP_KEEP if k in gap}
        slim_gap["reported"] = [
            {"id": m.get("id"), "name": m.get("name"), "value": m.get("value", "")}
            for m in gap.get("reported", [])
        ]
        slim_gap["missing"] = [{"id": m.get("id"), "name": m.get("name")} for m in gap.get("missing", [])]
        out["gap_analysis"] = slim_gap
    fd = out.get("five_dimensions")
    if isinstance(fd, dict):
        fd.pop("overall_provenance", None)
        for name in DIMENSIONS:
            if isinstance(fd.get(name), dict):
                fd[name].pop("provenance", None)
    # Methodology 2.0 blocks: keep the numbers, drop factor tables and quotes.
    ei = out.get("expected_impact")
    if isinstance(ei, dict):
        ei["outcomes"] = [{k: o.get(k) for k in ("kind", "unit", "stakeholder", "p10", "p50", "p90",
                                                  "uncertainty", "evidence_level")} for o in ei.get("outcomes", [])]
        ei.pop("targets", None)
    neg = out.get("negative_impacts")
    if isinstance(neg, dict):
        out["negative_impacts"] = {"material_count": neg.get("material_count"), "items": [
            {k: r.get(k) for k in ("id", "score", "band", "controlled")} for r in neg.get("items", [])]}
    gw = out.get("greenwashing")
    if isinstance(gw, dict) and isinstance(gw.get("claims_review"), dict):
        review = gw["claims_review"]
        gw["claims_review"] = {"score": review.get("score"), "flagged": review.get("flagged"),
                               "top": [{"risk": c.get("risk"), "text": str(c.get("text", ""))[:120]}
                                       for c in review.get("claims", [])[:3]]}
    for claim in out.get("impact_claims", []) or []:
        if isinstance(claim, dict):
            claim.pop("entities", None)
    return out


def _stamp(data: dict) -> dict[str, str]:
    from impact_vision.impact.methodology import methodology_stamp

    return data.get("methodology") or methodology_stamp()


def methodology_rows() -> list[tuple[str, str]]:
    """Methodology sheet: fixed evidence/standards notes + rows generated from the YAML."""
    from impact_vision.impact.methodology import methodology_appendix

    fixed = {topic: text for topic, text in METHODOLOGY}
    return [
        ("Evidence", fixed["Evidence"]),
        *methodology_appendix(),
        ("Gap analysis", fixed["Gap analysis"]),
        ("IC gate", fixed["IC gate"]),
        ("Standards", fixed["Standards"]),
        ("Schema", fixed["Schema"]),
    ]


def _ai(data: dict):  # type: ignore[no-untyped-def]
    from impact_vision.impact.ai_provenance import ai_provenance_for_report

    return ai_provenance_for_report(data)


def to_json(data: dict, *, slim: bool = False, indent: int | None = 2) -> str:
    payload = slim_report(data) if slim else dict(data)
    payload = {"schema_version": SCHEMA_VERSION, "export_mode": "slim" if slim else "full", **payload}
    payload["ai_provenance"] = _ai(data).model_dump(mode="json")
    payload["methodology"] = _stamp(data)
    return json.dumps(payload, indent=indent, default=str, ensure_ascii=False)


# ---------------------------------------------------------------- CSV

CSV_HEADER = ["Section", "Metric", "Value", "Details", "Numeric", "Max", "Unit"]


def csv_rows(data: dict) -> list[list[Any]]:
    """Rows shared by the CSV export and the XLSX "All figures" sheet."""
    company = data.get("company", {}) or {}
    rows: list[list[Any]] = [
        ["Company", "Name", company.get("name", ""), "", None, None, ""],
        ["Company", "Generated", data.get("generated_at", ""), "", None, None, ""],
        ["Company", "AI disclosure", _ai(data).disclosure, "", None, None, ""],
        ["Company", "Methodology", _stamp(data)["methodology_version"], f"config {_stamp(data)['config_hash']}",
         None, None, ""],
    ]
    fd = data.get("five_dimensions")
    if fd:
        rows.append(["5D", "Overall Grade", fd.get("overall_grade", ""),
                     f"{fd.get('overall_score')}/5.0", fd.get("overall_score"), 5, "score"])
        for name in DIMENSIONS:
            dim = fd.get(name) or {}
            if dim:
                rows.append(["5D", dim.get("dimension", name), f"{dim.get('score')}/5.0",
                             dim.get("notes", ""), dim.get("score"), 5, "score"])
    for a in data.get("sdg_alignments", []) or []:
        if a.get("score", 0) > 0:
            rows.append(["SDG", f"SDG {a['goal']}", f"{a['score']}/100",
                         f"{a.get('confidence', '')} | metrics: {','.join(a.get('matched_metrics', [])[:3])}",
                         a["score"], 100, "score"])
    gw = data.get("greenwashing") or {}
    if "overall_score" in gw:
        rows.append(["Greenwashing", "Risk score", f"{gw['overall_score']}/100",
                     gw.get("classification", ""), gw["overall_score"], 100, "risk"])
    ga = data.get("gap_analysis")
    if ga:
        rows.append(["Gap", "Coverage", f"{ga.get('coverage_percentage')}%", "",
                     ga.get("coverage_percentage"), 100, "%"])
        for m in ga.get("missing", []):
            rows.append(["Gap", m["id"], "MISSING", m.get("name", ""), None, None, ""])
        for m in ga.get("reported", []):
            number, unit = parse_number(m.get("value", ""))
            rows.append(["Gap", m["id"], m.get("value", ""), m.get("name", ""), number, None, unit])
    for claim in data.get("impact_claims", []) or []:
        rows.append(["Claim", claim.get("category", "intent"), claim.get("text", ""),
                     "metrics: " + ",".join(str(m) for m in claim.get("mapped_metrics", [])),
                     None, None, ""])
    return rows


def to_csv(data: dict) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(CSV_HEADER)
    for row in csv_rows(data):
        writer.writerow(["" if cell is None else cell for cell in row])
    return buf.getvalue()


# ---------------------------------------------------------------- XLSX

METHODOLOGY = (
    ("Evidence", "Claims are extracted from the documents and graded on the NESTA standards of "
                 "evidence (1–5); quantities are mapped to IRIS+ metric IDs only when unit and "
                 "context are unambiguous. Forward-looking commitments are never mapped."),
    ("5 Dimensions", "Scores (1–5) for What, Who, How Much, Contribution and Risk come from "
                     "reported metrics where available and from document text otherwise."),
    ("SDGs", "Scores (0–100) combine goal-specific metric coverage and inference from the "
             "business and impact themes; cross-cutting metrics count at reduced weight. "
             "'Material' marks goals the business model directly drives."),
    ("Gap analysis", "Coverage is the share of the sector's IRIS+ core metric set the company "
                     "reports."),
    ("Greenwashing", "Risk (0–100) weights claim–metric gaps (30%), missing negative impacts "
                     "(20%), vague language (20%), selective reporting (15%) and lack of "
                     "verification (15%). 60+ is flagged for review."),
    ("IC gate", "Applies the fund thesis thresholds. Failures caused by missing data are "
                "reported as insufficient evidence, not as a negative finding."),
    ("Standards", "IRIS+ 5.3c catalog, IMP Five Dimensions, UN SDG framework (17 goals, "
                  "169 targets)."),
    ("Schema", f"Export schema version {SCHEMA_VERSION}. Numeric cells are numbers; "
               "'—' marks values that were not reported."),
)


def _table(ws, headers: list[str], rows: list[list[Any]], *, widths: dict[int, float] | None = None,
           formats: dict[int, str] | None = None, start_row: int = 1) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F3A5F", end_color="1F3A5F", fill_type="solid")
    for col, title in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col, value=title)
        cell.font, cell.fill = header_font, header_fill
        cell.alignment = Alignment(vertical="center")
    for r, row in enumerate(rows, start_row + 1):
        for col, value in enumerate(row, 1):
            cell = ws.cell(row=r, column=col, value=value)
            if formats and col in formats and isinstance(value, (int, float)):
                cell.number_format = formats[col]
            if isinstance(value, str) and len(value) > 60:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    last_col = get_column_letter(len(headers))
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    ws.auto_filter.ref = f"A{start_row}:{last_col}{start_row + max(len(rows), 1)}"
    for col, title in enumerate(headers, 1):
        longest = max([len(str(title))] + [len(str(row[col - 1] or "")) for row in rows[:500]])
        width = (widths or {}).get(col) or min(max(10, longest + 2), 60)
        ws.column_dimensions[get_column_letter(col)].width = width


def build_workbook(data: dict):  # type: ignore[no-untyped-def]
    """Return an ``openpyxl.Workbook`` for *data* (raises ImportError without openpyxl)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    from impact_vision.impact.ai_provenance import ai_provenance_for_report, machine_marking, mark_xlsx

    mark_xlsx(wb, machine_marking(ai_provenance_for_report(data)))
    company = data.get("company", {}) or {}
    decision = data.get("decision") or {}
    fd = data.get("five_dimensions") or {}
    gw = data.get("greenwashing") or {}
    ga = data.get("gap_analysis") or {}

    ws = wb.active
    ws.title = "Summary"
    summary = [
        ["Company", company.get("name", "")],
        ["Sector", company.get("sector", "")],
        ["Geography", company.get("geography", "")],
        ["Generated", data.get("generated_at", "")],
        ["Standard", data.get("catalog_version", "")],
        ["IC gate", decision.get("display_status") or (decision.get("overall_status") or "").upper()],
        ["Recommendation", decision.get("recommendation", "")],
        ["5D overall (1–5)", _num(fd.get("overall_score"))],
        ["5D grade", fd.get("overall_grade", "")],
        ["Greenwashing risk (0–100)", _num(gw.get("overall_score"))],
        ["Core metric coverage (%)", _num(ga.get("coverage_percentage"))],
        ["Claims extracted", len(data.get("impact_claims", []) or [])],
        ["AI disclosure", _ai(data).disclosure],
        ["Methodology version", _stamp(data)["methodology_version"]],
        ["Methodology config hash", _stamp(data)["config_hash"]],
        ["Export schema", SCHEMA_VERSION],
    ]
    ws.append(["Impact assessment — " + str(company.get("name", ""))])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    _table(ws, ["Field", "Value"], summary, widths={1: 28, 2: 70}, start_row=3,
           formats={2: "0.0"})

    if fd:
        rows = []
        for name in DIMENSIONS:
            dim = fd.get(name) or {}
            if dim:
                rows.append([dim.get("dimension", name), _num(dim.get("score")),
                             dim.get("metrics_reported"), dim.get("metrics_available"),
                             dim.get("notes", "")])
        rows.append(["Overall", _num(fd.get("overall_score")), None, None, fd.get("overall_grade", "")])
        _table(wb.create_sheet("5 Dimensions"),
               ["Dimension", "Score (1–5)", "Metrics reported", "Metrics available", "Notes"],
               rows, widths={5: 70}, formats={2: "0.0"})

    alignments = [a for a in data.get("sdg_alignments", []) or [] if a.get("score", 0) > 0]
    if alignments:
        rows = [[a["goal"], a.get("goal_name", ""), _num(a["score"]), a.get("confidence", ""),
                 "yes" if a.get("material") else "no",
                 ", ".join(a.get("matched_metrics", [])[:8]), ", ".join(a.get("matched_targets", [])[:8])]
                for a in alignments]
        _table(wb.create_sheet("SDG Alignment"),
               ["SDG", "Goal", "Score (0–100)", "Confidence", "Material", "Matched metrics",
                "Matched targets"], rows, formats={3: "0.0"})

    if ga:
        rows = []
        for m in ga.get("reported", []):
            number, unit = parse_number(m.get("value", ""))
            rows.append(["Reported", m["id"], m.get("name", ""), number, unit, str(m.get("value", ""))])
        for m in ga.get("missing", []):
            rows.append(["Missing", m["id"], m.get("name", ""), None, "", "—"])
        _table(wb.create_sheet("Gap Analysis"),
               ["Status", "IRIS+ ID", "Metric", "Value", "Unit", "As reported"], rows,
               widths={3: 48}, formats={4: "#,##0.##"})

    claims = data.get("impact_claims", []) or []
    if claims:
        rows = [[c.get("category", ""), c.get("text", ""), _num(c.get("evidence_strength", c.get("evidence_level"))),
                 _num(c.get("confidence")), ", ".join(str(m) for m in c.get("mapped_metrics", [])),
                 c.get("extracted_by", "")]
                for c in claims]
        _table(wb.create_sheet("Claims"),
               ["Category", "Claim", "NESTA level", "Confidence", "Mapped metrics", "Extracted by"], rows,
               widths={2: 80}, formats={3: "0", 4: "0.00"})

    _table(wb.create_sheet("All figures"), CSV_HEADER, csv_rows(data), widths={3: 40, 4: 60},
           formats={5: "#,##0.##", 6: "0"})
    prov = _ai(data)
    ws_ai = wb.create_sheet("AI provenance")
    _table(ws_ai, ["Field", "Value"], [list(r) for r in prov.as_rows()], widths={1: 28, 2: 100})
    _table(ws_ai, ["Figure", "Stage", "Method", "Estimated", "Model", "Note"],
           [[r.figure, r.stage, r.method, "yes" if r.estimated else "no", r.model, r.note]
            for r in prov.records],
           widths={1: 34, 6: 40}, start_row=len(prov.as_rows()) + 3)
    _table(wb.create_sheet("Methodology"), ["Topic", "How it is calculated"],
           [list(row) for row in methodology_rows()] + [["AI use", prov.disclosure]], widths={1: 18, 2: 100})
    return wb


def write_xlsx(data: dict, path: str | Path) -> Path:
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(data).save(str(out))
    return out


__all__ = [
    "CSV_HEADER",
    "SCHEMA_VERSION",
    "build_workbook",
    "csv_rows",
    "parse_number",
    "slim_report",
    "to_csv",
    "to_json",
    "write_xlsx",
]
